from io import BytesIO
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.dependencies import get_current_user
from app.database import model_registry  # noqa: F401
from app.database import session as session_module
from app.database.base import Base
from app.modules.equipment.models import Equipment
from app.modules.equipment_types.models import EquipmentBrand, EquipmentModel, EquipmentType
from app.modules.batteries.models import Battery
from app.modules.faults_repairs.models import Fault
from app.modules.fuel.model_rate_import import MODEL_RATE_HEADERS
from app.modules.fuel.models import FuelModelConsumptionRate, FuelRecord
from app.modules.meter_readings.models import MeterReading
from app.modules.users.models import User
from web import main


def test_fuel_views_and_model_rate_excel_import(monkeypatch):
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session = sessionmaker(bind=engine, autoflush=False)
    Base.metadata.create_all(engine)

    db = testing_session()
    user = User(username="fuel-view-test", full_name="اختبار", hashed_password="unused", role="ADMIN")
    equipment_type = EquipmentType(name="نوع اختبار الوقود", measurement_unit="km")
    brand = EquipmentBrand(name="علامة اختبار الوقود")
    db.add_all([user, equipment_type, brand])
    db.flush()
    model = EquipmentModel(name="طراز اختبار الوقود", equipment_type_id=equipment_type.id, brand_id=brand.id)
    db.add(model)
    db.flush()
    equipment = Equipment(
        asset_code="FUEL-VIEW-1",
        registration_number="FUEL-VIEW-1",
        equipment_type_id=equipment_type.id,
        equipment_model_id=model.id,
        operational_status="available",
    )
    db.add(equipment)
    db.flush()
    fault = Fault(
        equipment_id=equipment.id,
        reported_date=date.today(),
        description="عطل بطارية لاختبار طلب الاستلام",
        severity="low",
    )
    db.add(fault)
    db.commit()
    user_id = user.id
    equipment_id = equipment.id
    model_id = model.id
    fault_id = fault.id
    db.close()

    def override_get_db():
        session = testing_session()
        try:
            yield session
        finally:
            session.close()

    def override_current_user():
        session = testing_session()
        try:
            yield session.get(User, user_id)
        finally:
            session.close()

    with monkeypatch.context() as patch:
        patch.setattr(main, "init_db", lambda: None)
        patch.setattr(main, "create_default_admin", lambda: None)
        patch.setattr(session_module, "SessionLocal", testing_session)
        app = main.create_app()
        app.dependency_overrides[session_module.get_db] = override_get_db
        app.dependency_overrides[get_current_user] = override_current_user
        with TestClient(app) as client:
            battery_page = client.get("/batteries")
            assert battery_page.status_code == 200
            assert 'id="batteryStatDialog"' in battery_page.text
            assert 'data-battery-stat="installed"' in battery_page.text
            assert 'data-battery-stat="total"' not in battery_page.text
            assert "الإجمالي · عرض الكل" not in battery_page.text
            assert "رقم الطلب" in battery_page.text
            assert "رقم تسلسلي غير مطلوب. مدة الاستبدال" not in battery_page.text
            assert "href=\"/spare-parts-received\"" not in battery_page.text
            assert "لا توجد بطاريات مسجلة حاليًا." in battery_page.text
            created_battery = client.post("/batteries", data={
                "battery_type": "AGM",
                "brand": "اختبار",
                "capacity": "60 Ah",
                "acquisition_request_item_id": client.post("/api/spare-parts-requests", json={
                    "request_number": "BAT-RECEIPT-1",
                    "request_date": date.today().isoformat(),
                    "received_date": date.today().isoformat(),
                    "recipient": "مستلم الاختبار",
                    "supplier_institution": "مورد الاختبار",
                    "source_type": "fault",
                    "source_id": fault_id,
                    "items": [{"asset_type": "battery", "part_name": "بطارية اختبار", "requested_quantity": 1, "received_quantity": 1}],
                }).json()["items"][0]["id"],
            }, follow_redirects=False)
            assert created_battery.status_code == 303
            verify = testing_session()
            battery_id = verify.query(Battery.id).one()[0]
            verify.close()
            battery_edit = client.get(f"/batteries/{battery_id}?edit=1")
            assert battery_edit.status_code == 200
            assert "تعديل بيانات البطارية" in battery_edit.text
            changed_battery = client.post(f"/batteries/{battery_id}/update", data={
                "battery_type": "AGM مطورة",
                "brand": "علامة معدلة",
                "capacity": "70 Ah",
                "manufacture_date": date.today().isoformat(),
                "notes": "ملاحظة اختبار",
            }, follow_redirects=False)
            assert changed_battery.status_code == 303
            battery_register = client.get("/batteries")
            assert battery_register.status_code == 200
            assert f"/batteries/{battery_id}?edit=1#battery-edit" in battery_register.text
            assert "تحديد للحذف" in battery_register.text
            assert '<th>الإجراءات</th>' not in battery_register.text
            assert 'class="battery-date-actions"' in battery_register.text
            deleted_battery = client.post(
                "/batteries/delete-selected",
                data={"battery_ids": [str(battery_id)]},
                follow_redirects=False,
            )
            assert deleted_battery.status_code == 303
            verify = testing_session()
            assert verify.query(Battery).count() == 0
            verify.close()
            empty_indicator_table = client.get("/batteries?filter=installed")
            assert empty_indicator_table.status_code == 200
            assert "لا توجد بطاريات مسجلة حاليًا." in empty_indicator_table.text

            fuel_page = client.get("/fuel")
            assert fuel_page.status_code == 200
            assert 'name="model_id"' in fuel_page.text
            assert 'name="equipment_id"' in fuel_page.text
            assert '<select id="fuelType" name="fuel_type" required>' in fuel_page.text
            assert "الطراز" in fuel_page.text and "رقم التسجيل" in fuel_page.text

            form_data = {
                "model_id": str(model_id + 1),
                "equipment_id": str(equipment_id),
                "fueling_date": date.today().isoformat(),
                "meter_value": "1000",
                "quantity": "20",
                "fuel_type": "ديزل",
            }
            mismatched_model = client.post("/fuel", data=form_data)
            assert mismatched_model.status_code == 400
            form_data["model_id"] = str(model_id)
            form_data["fuel_type"] = "نوع غير معتمد"
            invalid_fuel_type = client.post("/fuel", data=form_data)
            assert invalid_fuel_type.status_code == 400
            form_data["fuel_type"] = "ديزل"
            saved_fuel = client.post("/fuel", data=form_data, follow_redirects=False)
            assert saved_fuel.status_code == 303
            fuel_page = client.get("/fuel")
            assert '<th>الإجراءات</th>' not in fuel_page.text
            assert 'class="fuel-date-actions"' in fuel_page.text
            verify = testing_session()
            fuel_record = verify.query(FuelRecord).one()
            fuel_record_id = fuel_record.id
            verify.close()
            edit_page = client.get(f"/fuel?edit={fuel_record_id}")
            assert edit_page.status_code == 200
            assert 'id="fuelEditDialog"' in edit_page.text
            assert "تحديد للحذف" in edit_page.text
            changed_fuel = client.post(f"/fuel/{fuel_record_id}/update", data={
                "fueling_date": date.today().isoformat(),
                "meter_value": "1000",
                "quantity": "24.5",
                "fuel_type": "ديزل",
                "document_number": "FUEL-DOC-EDIT",
                "station": "محطة اختبار",
                "notes": "تعديل اختبار",
                "return_to": "/equipment/1/fuel",
            }, follow_redirects=False)
            assert changed_fuel.status_code == 303
            assert changed_fuel.headers["location"] == f"/equipment/{equipment_id}/fuel"
            equipment_history = client.get(f"/equipment/{equipment_id}/fuel")
            assert equipment_history.status_code == 200
            assert '<th>الإجراءات</th>' not in equipment_history.text
            assert 'class="fuel-date-actions"' in equipment_history.text
            assert f"/fuel?edit={fuel_record_id}" in equipment_history.text
            verify = testing_session()
            assert verify.get(FuelRecord, fuel_record_id).quantity == 24.5
            assert verify.query(MeterReading).filter_by(
                source_reference=f"fuel:{fuel_record_id}:meter"
            ).count() == 1
            verify.close()
            deleted_fuel = client.post(
                "/fuel/delete-selected",
                data={"record_ids": [str(fuel_record_id)], "return_to": f"/equipment/{equipment_id}/fuel"},
                follow_redirects=False,
            )
            assert deleted_fuel.status_code == 303
            assert deleted_fuel.headers["location"] == f"/equipment/{equipment_id}/fuel"
            verify = testing_session()
            assert verify.query(FuelRecord).count() == 0
            retained_meter = verify.query(MeterReading).filter_by(
                source_reference=None,
                equipment_id=equipment_id,
            ).one()
            assert retained_meter.odometer == 1000
            verify.close()

            rates_page = client.get("/fuel/model-rates")
            assert rates_page.status_code == 200
            assert 'id="fuelRateDialog"' in rates_page.text
            assert 'id="openFuelRateDialog"' in rates_page.text
            assert "تنزيل قالب Excel" in rates_page.text
            assert 'name="file"' in rates_page.text
            assert "العلامة التجارية" not in rates_page.text
            assert "الإجمالي · عرض الكل" not in rates_page.text

            added_rate = client.post("/fuel/model-rates", data={
                "model_id": str(model_id),
                "fuel_type": "ديزل",
                "consumption_rate": "8.0",
                "measurement_unit": "لتر/100 كم",
            }, follow_redirects=True)
            assert added_rate.status_code == 200
            assert "أُضيف معدل الاستهلاك بنجاح." in added_rate.text
            assert "8.00" in added_rate.text or ">8<" in added_rate.text

            rates_template = client.get("/fuel/model-rates/template.xlsx")
            assert rates_template.status_code == 200
            workbook = load_workbook(BytesIO(rates_template.content), read_only=True)
            sheet = workbook["معدلات الاستهلاك"]
            assert tuple(cell.value for cell in next(sheet.iter_rows())) == MODEL_RATE_HEADERS
            workbook.close()
            workbook = load_workbook(BytesIO(rates_template.content))
            sheet = workbook["معدلات الاستهلاك"]
            assert sheet.freeze_panes == "A2"
            assert sheet.data_validations.count == 3
            assert workbook.defined_names.get("FuelRateModels") is not None
            assert workbook.defined_names.get("FuelTypes") is not None
            assert "علامة اختبار الوقود" not in [value for row in workbook["قوائم"].iter_rows(values_only=True) for value in row]
            workbook.close()

            imported = BytesIO()
            workbook = load_workbook(BytesIO(rates_template.content))
            sheet = workbook["معدلات الاستهلاك"]
            sheet.append(("طراز اختبار الوقود", "ديزل", 12.5, "لتر/100 كم"))
            workbook.save(imported)
            workbook.close()
            response = client.post(
                "/fuel/model-rates/import",
                files={"file": ("fuel-rates.xlsx", imported.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                follow_redirects=True,
            )
            assert response.status_code == 200
            assert "تم استيراد أو تحديث 1 معدل استهلاك" in response.text
            assert "12.50" in response.text or "12.5" in response.text

            equipment_list = client.get("/fuel/equipment")
            assert equipment_list.status_code == 200
            assert "FUEL-VIEW-1" in equipment_list.text
            assert "/equipment/1/fuel" in equipment_list.text
            assert "<th>السجل</th>" not in equipment_list.text
            equipment_history = client.get(f"/equipment/{equipment_id}/fuel")
            assert equipment_history.status_code == 200

        verify = testing_session()
        assert verify.query(FuelModelConsumptionRate).count() == 1
        verify.close()
        app.dependency_overrides.clear()
    engine.dispose()
