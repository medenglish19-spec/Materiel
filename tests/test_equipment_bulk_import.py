from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database.base import Base
import app.database.model_registry  # noqa: F401
from app.modules.equipment import bulk_import
from app.modules.equipment.models import Equipment
from app.modules.equipment_types.models import EquipmentModel, EquipmentType


def _session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)(), engine


def _file(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "بيانات العتاد"
    sheet.append(bulk_import.HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parse_workbook_reads_template_columns_and_dates():
    parsed = bulk_import.parse_workbook(_file([("طراز", "R-1", "VIN-1", date(2024, 1, 2), None, "وثيقة")]))
    assert parsed[0]["row_number"] == 2
    assert parsed[0]["رقم التسجيل"] == "R-1"
    assert parsed[0]["تاريخ أول سير"].date() == date(2024, 1, 2)


def test_downloadable_template_has_approved_model_dropdown():
    db, engine = _session()
    try:
        equipment_type = EquipmentType(name="نوع قالب الاستيراد", measurement_unit="km")
        db.add(equipment_type)
        db.flush()
        db.add(EquipmentModel(name="طراز القالب", equipment_type_id=equipment_type.id))
        db.commit()

        workbook = load_workbook(BytesIO(bulk_import.build_template(db)))
        sheet = workbook["بيانات العتاد"]
        assert tuple(cell.value for cell in sheet[1]) == bulk_import.HEADERS
        assert sheet["A2"].value is None
        assert sheet.data_validations.dataValidation[0].formula1 == "=ApprovedEquipmentModels"
        assert "A2:A1001" in str(sheet.data_validations.dataValidation[0].sqref)
        assert workbook.defined_names["ApprovedEquipmentModels"].attr_text == "'الطرازات المعتمدة'!$A$2:$A$2"
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_import_keeps_valid_rows_and_reports_rejected_row_reason():
    db, engine = _session()
    try:
        equipment_type = EquipmentType(name="شاحنة اختبار الاستيراد", measurement_unit="km")
        db.add(equipment_type)
        db.flush()
        model = EquipmentModel(name="طراز اختبار الاستيراد", equipment_type_id=equipment_type.id)
        db.add(model)
        db.commit()

        template = load_workbook(BytesIO(bulk_import.build_template(db)))
        assert template["الطرازات المعتمدة"]["A2"].value == model.name
        rows = bulk_import.parse_workbook(_file([
            (model.name, "BULK-1", "BULK-VIN-1", date(2024, 1, 2), None, None),
            ("طراز اختبار الاستيراد مكتوب خطأ", "BULK-2", "BULK-VIN-2", date(2024, 1, 2), None, None),
        ]))
        preview = bulk_import.preview(db, rows)
        assert [row["ok"] for row in preview] == [True, False]
        assert "الطراز غير موجود" in preview[1]["errors"][0]
        result = bulk_import.import_rows(db, rows, user_id=None)

        assert [row["row_number"] for row in result["created"]] == [2]
        assert result["rejected"][0]["row_number"] == 3
        assert "الطراز غير موجود" in result["rejected"][0]["errors"][0]
        assert db.query(Equipment).count() == 1
        assert db.query(Equipment).filter_by(registration_number="BULK-1").one().first_service_date == date(2024, 1, 2)
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_import_rejects_identifier_already_in_database_without_stopping_other_rows():
    db, engine = _session()
    try:
        equipment_type = EquipmentType(name="نوع التكرار", measurement_unit="km")
        db.add(equipment_type)
        db.flush()
        model = EquipmentModel(name="طراز التكرار", equipment_type_id=equipment_type.id)
        db.add(model)
        db.flush()
        db.add(Equipment(asset_code="EXISTING-BULK", registration_number="DUP-1", equipment_type_id=equipment_type.id, equipment_model_id=model.id))
        db.commit()
        rows = bulk_import.parse_workbook(_file([
            (model.name, "DUP-1", "VIN-DUP", date(2024, 1, 2), None, None),
            (model.name, "NEW-1", "VIN-NEW", date(2024, 1, 2), None, None),
        ]))
        result = bulk_import.import_rows(db, rows, user_id=None)
        assert len(result["created"]) == 1
        assert len(result["rejected"]) == 1
        assert "رقم التسجيل مستخدم" in result["rejected"][0]["errors"][0]
        assert db.query(Equipment).count() == 2
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_preview_rejects_model_name_that_resolves_to_multiple_models():
    db, engine = _session()
    try:
        first_type = EquipmentType(name="نوع أول للاسم الملتبس", measurement_unit="km")
        second_type = EquipmentType(name="نوع ثان للاسم الملتبس", measurement_unit="km")
        db.add_all([first_type, second_type])
        db.flush()
        db.add_all([
            EquipmentModel(name="طراز مشترك", equipment_type_id=first_type.id),
            EquipmentModel(name="طراز مشترك", equipment_type_id=second_type.id),
        ])
        db.commit()
        rows = bulk_import.parse_workbook(_file([("طراز مشترك", "AMB-1", "AMB-VIN", date(2024, 1, 2), None, None)]))
        result = bulk_import.preview(db, rows)
        assert result[0]["ok"] is False
        assert "ملتبس" in result[0]["errors"][0]
    finally:
        db.close()
        Base.metadata.drop_all(engine)
