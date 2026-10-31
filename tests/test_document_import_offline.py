from pathlib import Path
import io
import sys
import types
import pytest


def test_document_import_module_exists():
    assert Path("app/modules/document_import").exists()
    assert Path("app/modules/document_import/router.py").exists()
    assert Path("app/modules/document_import/services.py").exists()
    assert Path("app/modules/document_import/extraction/extractor.py").exists()
    assert Path("app/modules/document_import/parsing/parser.py").exists()
    assert Path("app/modules/document_import/matching/matcher.py").exists()
    assert Path("app/modules/document_import/normalization/normalizer.py").exists()


def test_router_has_endpoints():
    s = Path("app/modules/document_import/router.py").read_text(encoding="utf-8", errors="replace")
    assert "/document-import/preview" in s
    assert "/document-import/apply" in s


def test_services_have_real_logic():
    s = Path("app/modules/document_import/services.py").read_text(encoding="utf-8", errors="replace")
    assert "preview_from_document" in s
    assert "apply_candidates" in s
    assert "extract_document" in s
    assert "match_against_definitions" in s


def test_extractor_offline_refs():
    s = Path("app/modules/document_import/extraction/extractor.py").read_text(encoding="utf-8", errors="replace")
    assert "docling" in s.lower() or "pypdf" in s.lower()


def test_schemas_have_review_fields():
    s = Path("app/modules/document_import/schemas.py").read_text(encoding="utf-8", errors="replace")
    assert "approved" in s
    assert "ignored" in s
    assert "edited" in s
    assert "ImportApplyRequest" in s
    assert "ImportApplyItem" in s


def test_document_import_forms_join_the_shared_edit_manager():
    template = Path("app/modules/document_import/templates/upload.html").read_text(encoding="utf-8")
    assert 'id="previewForm"' in template and 'id="previewForm" action="/document-import/preview" method="post" enctype="multipart/form-data" data-em-root' in template
    assert 'id="applyForm"' in template and 'id="applyForm" method="post" action="/document-import/apply" style="margin-top:16px" data-em-root' in template


def test_empty_xlsx_returns_clear_arabic_extraction_error(tmp_path):
    from openpyxl import Workbook

    from app.modules.document_import.extraction.extractor import ExtractionError, extract_document

    path = tmp_path / "empty.xlsx"
    Workbook().save(path)

    with pytest.raises(ExtractionError, match="لم يُعثر على بيانات"):
        extract_document(path)


def test_corrupt_xlsx_returns_clear_arabic_extraction_error(tmp_path):
    from app.modules.document_import.extraction.extractor import ExtractionError, extract_document

    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an xlsx file")

    with pytest.raises(ExtractionError, match="ملف XLSX تالف"):
        extract_document(path)


def test_xls_extraction_returns_nonempty_rows_and_skips_blank_rows(tmp_path, monkeypatch):
    from app.modules.document_import.extraction.extractor import extract_document

    class Sheet:
        name = "Equipment"
        nrows = 3
        ncols = 2

        def cell_value(self, row, col):
            return [["Model", "Capacity"], ["", ""], ["Truck", 12]][row][col]

    class Workbook:
        def sheets(self):
            return [Sheet()]

        def release_resources(self):
            pass

    monkeypatch.setitem(sys.modules, "xlrd", types.SimpleNamespace(open_workbook=lambda *args, **kwargs: Workbook()))
    result = extract_document(tmp_path / "legacy.xls")

    assert result.tables[0].rows == [["Model", "Capacity"], ["Truck", "12"]]
