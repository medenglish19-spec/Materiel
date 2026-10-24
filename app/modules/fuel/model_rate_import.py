"""Excel template and validation for fuel consumption rates by model."""

from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session, joinedload

from app.modules.equipment_types.models import EquipmentModel
from app.modules.fuel.models import FuelModelConsumptionRate


MODEL_RATE_HEADERS = ("العلامة التجارية", "الطراز", "نوع الوقود", "معدل الاستهلاك", "وحدة القياس")
FUEL_TYPE_OPTIONS = ("ديزل", "بنزين", "غاز بترولي مسال", "غاز طبيعي")
CONSUMPTION_UNIT_OPTIONS = ("لتر/100 كم", "لتر/ساعة")


def list_models(db: Session) -> list[EquipmentModel]:
    return (
        db.query(EquipmentModel)
        .options(joinedload(EquipmentModel.brand), joinedload(EquipmentModel.equipment_type))
        .order_by(EquipmentModel.name, EquipmentModel.id)
        .all()
    )


def model_label(model: EquipmentModel) -> str:
    brand = model.brand.name if model.brand else "بدون علامة تجارية"
    return f"{brand} · {model.name}"


def build_model_rate_template(models: list[EquipmentModel]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "معدلات الاستهلاك"
    sheet.append(MODEL_RATE_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E501"
    widths = (24, 28, 22, 20, 20)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    brands = sorted({model.brand.name for model in models if model.brand and model.brand.name})
    model_names = sorted({model.name for model in models if model.name})
    lists = workbook.create_sheet("قوائم")
    lists.sheet_state = "hidden"
    for row, value in enumerate(FUEL_TYPE_OPTIONS, start=1):
        lists.cell(row=row, column=1, value=value)
    for row, value in enumerate(CONSUMPTION_UNIT_OPTIONS, start=1):
        lists.cell(row=row, column=2, value=value)
    for row, value in enumerate(brands, start=1):
        lists.cell(row=row, column=3, value=value)
    for row, value in enumerate(model_names, start=1):
        lists.cell(row=row, column=4, value=value)

    workbook.defined_names.add(DefinedName(
        "FuelTypes", attr_text=f"'قوائم'!$A$1:$A${len(FUEL_TYPE_OPTIONS)}"
    ))
    workbook.defined_names.add(DefinedName(
        "FuelRateUnits", attr_text=f"'قوائم'!$B$1:$B${len(CONSUMPTION_UNIT_OPTIONS)}"
    ))

    fuel_validation = DataValidation(type="list", formula1=f"'قوائم'!$A$1:$A${len(FUEL_TYPE_OPTIONS)}", allow_blank=False)
    unit_validation = DataValidation(type="list", formula1=f"'قوائم'!$B$1:$B${len(CONSUMPTION_UNIT_OPTIONS)}", allow_blank=False)
    if brands:
        workbook.defined_names.add(DefinedName("FuelBrands", attr_text=f"'قوائم'!$C$1:$C${len(brands)}"))
        brand_validation = DataValidation(type="list", formula1=f"'قوائم'!$C$1:$C${len(brands)}", allow_blank=False)
        sheet.add_data_validation(brand_validation)
        brand_validation.add("A2:A501")
    if model_names:
        workbook.defined_names.add(DefinedName("FuelModels", attr_text=f"'قوائم'!$D$1:$D${len(model_names)}"))
        model_validation = DataValidation(type="list", formula1=f"'قوائم'!$D$1:$D${len(model_names)}", allow_blank=False)
        sheet.add_data_validation(model_validation)
        model_validation.add("B2:B501")
    fuel_validation.formula1 = "=FuelTypes"
    unit_validation.formula1 = "=FuelRateUnits"
    if brands:
        brand_validation.formula1 = "=FuelBrands"
    if model_names:
        model_validation.formula1 = "=FuelModels"
    for validation, target in ((fuel_validation, "C2:C501"), (unit_validation, "E2:E501")):
        sheet.add_data_validation(validation)
        validation.add(target)
    sheet["D2"].number_format = "0.00"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_model_rate_workbook(content: bytes, models: list[EquipmentModel]) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("تعذر قراءة الملف؛ استخدم قالب Excel بصيغة .xlsx.") from exc
    if "معدلات الاستهلاك" not in workbook.sheetnames:
        raise ValueError("يجب أن يحتوي الملف ورقة «معدلات الاستهلاك».")
    sheet = workbook["معدلات الاستهلاك"]
    headers = tuple(str(cell or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if headers != MODEL_RATE_HEADERS:
        raise ValueError("رأس الجدول غير مطابق لقالب معدلات استهلاك الوقود.")

    model_map = {
        ((model.brand.name if model.brand else "بدون علامة تجارية").strip().casefold(), model.name.strip().casefold()): model
        for model in models
    }
    rows: list[dict] = []
    seen: set[tuple[int, str]] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = tuple(values[:5])
        if not any(value not in (None, "") for value in values):
            continue
        brand, model_name, fuel_type, raw_rate, unit = (str(value or "").strip() for value in values)
        model = model_map.get((brand.casefold(), model_name.casefold()))
        if model is None:
            raise ValueError(f"العلامة التجارية والطراز في الصف {row_number} غير موجودين في بيانات البرنامج.")
        if fuel_type not in FUEL_TYPE_OPTIONS:
            raise ValueError(f"نوع الوقود في الصف {row_number} غير مدرج في القائمة المعتمدة.")
        if unit not in CONSUMPTION_UNIT_OPTIONS:
            raise ValueError(f"وحدة القياس في الصف {row_number} غير مدرجة في القائمة المعتمدة.")
        try:
            rate = Decimal(raw_rate)
        except (InvalidOperation, ValueError):
            raise ValueError(f"معدل الاستهلاك في الصف {row_number} يجب أن يكون رقمًا.") from None
        if not rate.is_finite() or rate <= 0:
            raise ValueError(f"معدل الاستهلاك في الصف {row_number} يجب أن يكون أكبر من صفر.")
        key = (model.id, fuel_type)
        if key in seen:
            raise ValueError(f"يوجد تكرار للطراز ونوع الوقود في الصف {row_number}.")
        seen.add(key)
        rows.append({
            "equipment_model_id": model.id,
            "fuel_type": fuel_type,
            "consumption_rate": rate,
            "measurement_unit": unit,
        })
    if not rows:
        raise ValueError("لا يحتوي الملف على معدلات استهلاك لإدخالها.")
    return rows


def save_model_rates(db: Session, rows: list[dict]) -> int:
    for row in rows:
        existing = db.query(FuelModelConsumptionRate).filter_by(
            equipment_model_id=row["equipment_model_id"], fuel_type=row["fuel_type"]
        ).one_or_none()
        if existing:
            existing.consumption_rate = row["consumption_rate"]
            existing.measurement_unit = row["measurement_unit"]
        else:
            db.add(FuelModelConsumptionRate(**row))
    db.commit()
    return len(rows)
