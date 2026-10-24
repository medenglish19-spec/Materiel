"""Excel template, parsing, validation, and row-wise equipment import."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.modules.equipment import services
from app.modules.equipment.schemas import EquipmentCreate
from app.modules.equipment.models import Equipment
from app.modules.equipment_types.models import EquipmentModel, EquipmentType

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
HEADERS = ("الطراز", "رقم التسجيل", "رقم الهيكل", "تاريخ أول سير", "تاريخ الاقتناء", "الوثيقة")


def build_template(db: Session) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "بيانات العتاد"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5B")
        cell.alignment = Alignment(horizontal="center")
    for column, width in enumerate((30, 22, 24, 20, 20, 28), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    for row_number in range(2, MAX_ROWS + 2):
        sheet.cell(row_number, 2).number_format = "@"
        sheet.cell(row_number, 3).number_format = "@"
        sheet.cell(row_number, 4).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 5).number_format = "yyyy-mm-dd"

    models_sheet = workbook.create_sheet("الطرازات المعتمدة")
    models_sheet.append(("الطراز",))
    models = (
        db.query(EquipmentModel)
        .join(EquipmentType, EquipmentModel.equipment_type_id == EquipmentType.id)
        .filter(EquipmentModel.is_frozen.is_(False), EquipmentType.is_frozen.is_(False))
        .order_by(EquipmentType.name, EquipmentModel.name)
        .all()
    )
    name_counts = {}
    for model in db.query(EquipmentModel.name).all():
        name_counts[model.name.strip().casefold()] = name_counts.get(model.name.strip().casefold(), 0) + 1
    approved_model_count = 0
    for model in models:
        if name_counts.get(model.name.strip().casefold()) == 1:
            models_sheet.append((model.name,))
            approved_model_count += 1
    for cell in models_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5B")
    models_sheet.column_dimensions["A"].width = 40
    if approved_model_count:
        workbook.defined_names.add(DefinedName(
            "ApprovedEquipmentModels",
            attr_text=f"'{models_sheet.title}'!$A$2:$A${approved_model_count + 1}",
        ))
        model_validation = DataValidation(
            type="list",
            formula1="=ApprovedEquipmentModels",
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="طراز غير معتمد",
            error="اختر الطراز من القائمة المنسدلة للطرازات المعتمدة.",
            showInputMessage=True,
            promptTitle="اختيار الطراز",
            prompt="اختر اسم الطراز المسجل في مركز البيانات.",
        )
        sheet.add_data_validation(model_validation)
        model_validation.add(f"A2:A{MAX_ROWS + 1}")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _date(value, label: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"{label}: استخدم صيغة YYYY-MM-DD") from exc


def parse_workbook(data: bytes) -> list[dict]:
    if not data:
        raise ValueError("الملف فارغ.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("حجم الملف يتجاوز الحد المسموح (5 ميغابايت).")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("تعذر قراءة الملف. ارفع قالب Excel بصيغة .xlsx.") from exc
    sheet = workbook["بيانات العتاد"] if "بيانات العتاد" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise ValueError("الملف لا يحتوي على صف العناوين.") from exc
    normalized = [str(value or "").replace("\u00a0", " ").strip() for value in header]
    indexes = {name: normalized.index(name) for name in HEADERS if name in normalized}
    required = ("الطراز",)
    missing = [name for name in required if name not in indexes]
    if missing:
        raise ValueError("عناوين الأعمدة المطلوبة غير موجودة: " + "، ".join(missing))
    parsed = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(parsed) >= MAX_ROWS:
            raise ValueError(f"يتجاوز الملف الحد الأقصى وهو {MAX_ROWS} صف.")
        parsed.append({"row_number": row_number, **{name: values[index] if index < len(values) else None for name, index in indexes.items()}})
    if not parsed:
        raise ValueError("لا توجد صفوف بيانات في الملف.")
    return parsed


def _prepare_rows(db: Session, rows: list[dict]) -> list[dict]:
    models = db.query(EquipmentModel).join(EquipmentType).all()
    model_index = {}
    for model in models:
        model_index.setdefault(model.name.strip().casefold(), []).append(model)
    existing_reg = {v.casefold() for (v,) in db.query(Equipment.registration_number).filter(Equipment.registration_number.is_not(None)).all()}
    existing_vin = {v.casefold() for (v,) in db.query(Equipment.vin).filter(Equipment.vin.is_not(None)).all()}
    seen_reg: dict[str, int] = {}
    seen_vin: dict[str, int] = {}
    prepared = []
    for row in rows:
        errors = []
        model_name = _text(row.get("الطراز"))
        candidates = model_index.get((model_name or "").casefold(), [])
        if not candidates:
            errors.append("الطراز غير موجود في البيانات الأساسية.")
        elif len(candidates) > 1:
            errors.append("اسم الطراز ملتبس؛ يوجد أكثر من طراز مسجل بهذا الاسم، لذلك تعذر تحديد النوع المرتبط به.")
        elif candidates[0].is_frozen:
            errors.append("طراز العتاد مجمّد.")
        elif candidates[0].equipment_type.is_frozen:
            errors.append("نوع العتاد المرتبط بالطراز مجمّد.")
        registration = _text(row.get("رقم التسجيل"))
        vin = _text(row.get("رقم الهيكل"))
        for value, label, known, seen in ((registration, "رقم التسجيل", existing_reg, seen_reg), (vin, "رقم الهيكل", existing_vin, seen_vin)):
            if value:
                folded = value.casefold()
                if folded in known:
                    errors.append(f"{label} مستخدم في سجل عتاد موجود.")
                elif folded in seen:
                    errors.append(f"{label} مكرر في الصف {seen[folded]}.")
                else:
                    seen[folded] = row["row_number"]
        try:
            first_service_date = _date(row.get("تاريخ أول سير"), "تاريخ أول سير")
            acquisition_date = _date(row.get("تاريخ الاقتناء"), "تاريخ الاقتناء")
        except ValueError as exc:
            errors.append(str(exc))
            first_service_date = acquisition_date = None
        if first_service_date is None:
            errors.append("تاريخ أول سير مطلوب.")
        data = None
        if len(candidates) == 1:
            try:
                data = EquipmentCreate(
                    equipment_type_id=candidates[0].equipment_type_id,
                    equipment_model_id=candidates[0].id,
                    registration_number=registration,
                    vin=vin,
                    first_service_date=first_service_date,
                    acquisition_date=acquisition_date,
                    acquisition_document=_text(row.get("الوثيقة")),
                )
            except ValidationError as exc:
                errors.append(f"البيانات غير صالحة: {exc}")
        prepared.append({"row_number": row["row_number"], "data": data, "errors": errors})
    return prepared


def preview(db: Session, rows: list[dict]) -> list[dict]:
    prepared = _prepare_rows(db, rows)
    return [{"row_number": row["row_number"], "ok": not row["errors"], "errors": row["errors"]} for row in prepared]


def import_rows(db: Session, rows: list[dict], user_id: int) -> dict:
    prepared = _prepare_rows(db, rows)
    created = []
    rejected = []
    for row in prepared:
        if row["errors"]:
            rejected.append({"row_number": row["row_number"], "errors": row["errors"]})
            continue
        try:
            item = services.create_equipment(db, row["data"], user_id=user_id)
            created.append({"row_number": row["row_number"], "equipment_id": item.id})
        except ValueError as exc:
            db.rollback()
            rejected.append({"row_number": row["row_number"], "errors": [str(exc) or "بيانات الصف غير صالحة."]})
        except IntegrityError:
            db.rollback()
            rejected.append({"row_number": row["row_number"], "errors": ["تعارض مع قيد فريد في قاعدة البيانات، مثل رقم تسجيل أو رقم هيكل مستخدم."]})
        except SQLAlchemyError:
            db.rollback()
            rejected.append({"row_number": row["row_number"], "errors": ["تعذر حفظ الصف بسبب خطأ في قاعدة البيانات."]})
    return {"created": created, "rejected": rejected}
