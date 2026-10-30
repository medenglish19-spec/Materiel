"""Excel template and direct importer for model-scoped maintenance operations."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import unicodedata
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.modules.equipment_types.models import EquipmentModel
from app.modules.maintenance.models import (
    MaintenanceModelOperation,
    MaintenanceOperation,
    MaintenanceOperationGroup,
)


GROUP_NO_VALUE = "بدون مجموعة"
MAX_OPERATION_IMPORT_ROWS = 2000
OPERATION_IMPORT_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

OPERATION_IMPORT_HEADERS = (
    "العملية",
    "الوصف",
    "المجموعة",
    "شرط 1 (كم)",
    "شرط 2 (ساعة)",
    "شرط 3 (يوم)",
    "إنذار 1 (كم)",
    "إنذار 2 (ساعة)",
    "إنذار 3 (يوم)",
)
GENERAL_OPERATION_IMPORT_HEADERS = (
    "العملية",
    "الوصف",
    "المجموعة",
    "إنذار كم",
    "إنذار ساعة تشغيل",
    "إنذار يوم",
)
def build_operation_import_template(*, group_names: list[str]) -> bytes:
    """Create a blank, validated workbook reusable for any model-operation import."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "العمليات"
    sheet.sheet_view.rightToLeft = True
    sheet.append(OPERATION_IMPORT_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{MAX_OPERATION_IMPORT_ROWS + 1}"

    header_fill = PatternFill("solid", fgColor="315F42")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32

    widths = (32, 42, 28, 18, 20, 16, 18, 20, 16)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width

    comments = (
        "مطلوب؛ لا يتجاوز 120 حرفًا، ويجب ألا يتكرر داخل الملف.",
        "اختياري.",
        "اختياري؛ اختر اسم مجموعة موجودة أو «بدون مجموعة».",
        "شرط بالكيلومتر؛ اختياري، ويجب أن تكون القيمة أكبر من صفر.",
        "شرط بساعات التشغيل؛ اختياري، ويجب أن تكون القيمة أكبر من صفر.",
        "شرط بالأيام؛ اختياري، أدخل عددًا صحيحًا أكبر من صفر.",
        "إنذار بالكيلومتر؛ اختياري، والقيمة صفر أو أكبر.",
        "إنذار بساعات التشغيل؛ اختياري، والقيمة صفر أو أكبر.",
        "إنذار بالأيام؛ اختياري، أدخل عددًا صحيحًا صفرًا أو أكبر.",
    )
    for column, instruction in enumerate(comments, start=1):
        sheet.cell(1, column).comment = Comment(instruction, "Materiel")

    group_sheet = workbook.create_sheet("القوائم")
    group_choices = ["بدون مجموعة", *group_names]
    for row_number, group_name in enumerate(group_choices, start=1):
        group_sheet.cell(row_number, 1, group_name)
    group_sheet.sheet_state = "hidden"
    workbook.defined_names.add(DefinedName(
        "MaintenanceGroupOptions",
        attr_text=f"'القوائم'!$A$1:$A${len(group_choices)}",
    ))
    groups_validation = DataValidation(type="list", formula1="=MaintenanceGroupOptions", allow_blank=True)
    groups_validation.error = "اختر مجموعة من القائمة أو اترك الحقل فارغًا."
    groups_validation.errorTitle = "مجموعة غير معروفة"
    groups_validation.errorStyle = "stop"
    groups_validation.showErrorMessage = True
    sheet.add_data_validation(groups_validation)
    groups_validation.add(f"C2:C{MAX_OPERATION_IMPORT_ROWS + 1}")

    for column in range(4, 10):
        is_warning = column >= 7
        is_day = column in {6, 9}
        positive_decimal = DataValidation(
            type="whole" if is_day else "decimal",
            operator="greaterThanOrEqual" if is_warning else "greaterThan",
            formula1="0", allow_blank=True,
        )
        positive_decimal.error = "أدخل قيمة إنذار صفرًا أو أكبر." if is_warning else "أدخل قيمة شرط أكبر من صفر."
        positive_decimal.errorTitle = "قيمة غير صالحة"
        positive_decimal.errorStyle = "stop"
        positive_decimal.showErrorMessage = True
        sheet.add_data_validation(positive_decimal)
        letter = chr(64 + column)
        positive_decimal.add(f"{letter}2:{letter}{MAX_OPERATION_IMPORT_ROWS + 1}")

    instructions = workbook.create_sheet("إرشادات")
    instructions.sheet_view.rightToLeft = True
    instructions.append(("قالب استيراد عمليات الصيانة",))
    instructions.append(("هذا قالب عام؛ يحدد المستخدم الطراز المستهدف عند رفع الملف إلى البرنامج.",))
    instructions.append(("أدخل كل عملية في صف مستقل بورقة «العمليات»، ولا تغيّر أسماء الأعمدة أو ترتيبها.",))
    instructions.append(("أدخل الشروط في أعمدة الكيلومترات وساعات التشغيل والأيام؛ اترك غير المستخدم فارغًا.",))
    instructions.append(("أدخل الإنذارات في أعمدتها؛ وحدات القياس ثابتة ومبينة في رأس كل عمود.",))
    instructions.append(("كل عملية تحتاج شرطًا واحدًا على الأقل وإنذارًا واحدًا على الأقل.",))
    instructions.append(("قيم الأيام أعداد صحيحة؛ شرط اليوم أكبر من صفر وإنذار اليوم صفر أو أكبر.",))
    instructions.append(("اسم المجموعة اختياري؛ اختر مجموعة من القائمة أو اترك الحقل فارغًا.",))
    instructions.column_dimensions["A"].width = 110
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    instructions["A1"].font = Font(name="Arial", bold=True, size=14, color="315F42")
    instructions.row_dimensions[1].height = 26

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_general_operation_import_template(*, group_names: list[str]) -> bytes:
    """Create the Excel template for the shared, condition-free operation library."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "العمليات"
    sheet.sheet_view.rightToLeft = True
    sheet.append(GENERAL_OPERATION_IMPORT_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{MAX_OPERATION_IMPORT_ROWS + 1}"

    header_fill = PatternFill("solid", fgColor="315F42")
    comments = (
        "مطلوب؛ لا يتجاوز 120 حرفًا.",
        "اختياري.",
        "اختياري؛ اختر مجموعة من القائمة أو «بدون مجموعة».",
        "أدخل قيمة إنذار واحدة على الأقل في أحد أعمدة الإنذار.",
        "أدخل قيمة إنذار واحدة على الأقل في أحد أعمدة الإنذار.",
        "اختياري؛ إذا أدخلت قيمة أيام فيجب أن تكون عددًا صحيحًا.",
    )
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.comment = Comment(comments[cell.column - 1], "Materiel")
    sheet.row_dimensions[1].height = 32
    for column, width in enumerate((32, 42, 28, 18, 20, 16), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width

    lists = workbook.create_sheet("القوائم")
    for row_number, value in enumerate([GROUP_NO_VALUE, *group_names], start=1):
        lists.cell(row_number, 1, value)
    lists.sheet_state = "hidden"
    workbook.defined_names.add(DefinedName(
        "MaintenanceGroupOptions",
        attr_text=f"'القوائم'!$A$1:$A${max(1, len(group_names) + 1)}",
    ))
    groups_validation = DataValidation(
        type="list", formula1="=MaintenanceGroupOptions", allow_blank=True,
    )
    groups_validation.error = "اختر مجموعة من القائمة أو اترك الحقل بلا مجموعة."
    groups_validation.errorTitle = "قيمة المجموعة غير صالحة"
    groups_validation.errorStyle = "stop"
    groups_validation.showErrorMessage = True
    sheet.add_data_validation(groups_validation)
    groups_validation.add(f"C2:C{MAX_OPERATION_IMPORT_ROWS + 1}")

    for column in (4, 5, 6):
        validation = DataValidation(
            type="whole" if column == 7 else "decimal",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=True,
        )
        validation.error = "أدخل قيمة إنذار صفرًا أو أكبر."
        validation.errorTitle = "قيمة الإنذار غير صالحة"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        letter = chr(64 + column)
        validation.add(f"{letter}2:{letter}{MAX_OPERATION_IMPORT_ROWS + 1}")

    instructions = workbook.create_sheet("إرشادات")
    instructions.sheet_view.rightToLeft = True
    for instruction in (
        "قالب استيراد عمليات مكتبة الصيانة العامة",
        "هذا القالب للمكتبة العامة المشتركة بين الطرازات، ولا يتضمن شروط استحقاق الخطط.",
        "أدخل كل عملية في صف مستقل بورقة «العمليات»، ولا تغيّر أسماء الأعمدة أو ترتيبها.",
        "أدخل قيمة إنذار واحدة على الأقل؛ الوحدات محددة في عناوين الأعمدة.",
        "قيم الإنذار صفر أو أكبر. قيمة الإنذار بالأيام عدد صحيح.",
        "تُضاف العمليات الجديدة إلى المكتبة العامة بحالة مفعلة.",
        "المجموعة اختيارية؛ اختر اسم مجموعة من القائمة أو «بدون مجموعة».",
        "يحدد شرط الاستحقاق عند استخدام العملية في خطة الطراز.",
    ):
        instructions.append((instruction,))
    instructions.column_dimensions["A"].width = 100
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    instructions["A1"].font = Font(name="Arial", bold=True, size=14, color="315F42")
    instructions.row_dimensions[1].height = 26

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _clean_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _identity(value: str) -> str:
    return unicodedata.normalize("NFKC", _clean_text(value)).casefold()


def _number(value, *, field: str, row_number: int) -> Decimal | None:
    if value is None or _clean_text(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"الصف {row_number}: أدخل رقمًا صالحًا في «{field}».")
    normalized = _clean_text(value).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))
    normalized = normalized.replace("٬", "").replace(",", "").replace("٫", ".")
    try:
        parsed = Decimal(normalized)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(f"الصف {row_number}: أدخل رقمًا صالحًا في «{field}».") from exc
    if not parsed.is_finite() or parsed <= 0:
        raise ValueError(f"الصف {row_number}: يجب أن تكون قيمة «{field}» أكبر من صفر.")
    return parsed


def _nonnegative_number(value, *, field: str, row_number: int) -> Decimal | None:
    if value is None or _clean_text(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"الصف {row_number}: أدخل رقمًا صالحًا في «{field}».")
    normalized = _clean_text(value).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))
    normalized = normalized.replace("٬", "").replace(",", "").replace("٫", ".")
    try:
        parsed = Decimal(normalized)
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(f"الصف {row_number}: أدخل رقمًا صالحًا في «{field}».") from exc
    if not parsed.is_finite() or parsed < 0:
        raise ValueError(f"الصف {row_number}: يجب ألا تقل قيمة «{field}» عن صفر.")
    return parsed


def _parse_fixed_axes(
    values: list,
    *,
    start_index: int,
    axes: tuple[tuple[str, str], ...],
    field_name: str,
    row_number: int,
    required: bool,
    warning: bool = False,
) -> dict:
    parsed = {}
    present = False
    parser = _nonnegative_number if warning else _number
    for offset, (field, unit_label) in enumerate(axes):
        value = parser(
            values[start_index + offset] if start_index + offset < len(values) else None,
            field=f"{field_name} ({unit_label})",
            row_number=row_number,
        )
        if value is not None and unit_label == "يوم" and value != value.to_integral_value():
            raise ValueError(f"الصف {row_number}: قيمة {field_name} بالأيام يجب أن تكون عددًا صحيحًا.")
        parsed[field] = int(value) if value is not None and unit_label == "يوم" else value
        present = present or value is not None

    if required and not present:
        raise ValueError(f"الصف {row_number}: أدخل قيمة {field_name} واحدة على الأقل.")
    return parsed


def parse_operation_workbook(data: bytes, groups: list[MaintenanceOperationGroup]) -> list[dict]:
    """Validate the expected workbook structure and return normalized rows."""
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, ValueError, zipfile.BadZipFile, OSError) as exc:
        raise ValueError("تعذر قراءة الملف. ارفع ملف Excel صالحًا بصيغة .xlsx.") from exc
    try:
        if "العمليات" not in workbook.sheetnames:
            raise ValueError("يجب أن يحتوي الملف على ورقة باسم «العمليات».")
        sheet = workbook["العمليات"]
        headers = tuple(_clean_text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        if headers != OPERATION_IMPORT_HEADERS:
            expected_headers = "، ".join(OPERATION_IMPORT_HEADERS)
            raise ValueError(f"رؤوس الأعمدة أو ترتيبها غير مطابق. استخدم الترتيب التالي: {expected_headers}.")

        group_lookup: dict[str, int] = {}
        ambiguous_groups: set[str] = set()
        for group in groups:
            key = _identity(group.name)
            if key in group_lookup and group_lookup[key] != group.id:
                ambiguous_groups.add(key)
            group_lookup[key] = group.id

        rows: list[dict] = []
        seen_names: set[str] = set()
        row_count = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=len(OPERATION_IMPORT_HEADERS), values_only=True), start=2):
            values = list(values)
            if not any(_clean_text(value) for value in values):
                continue
            row_count += 1
            if row_count > MAX_OPERATION_IMPORT_ROWS:
                raise ValueError(f"الحد الأقصى {MAX_OPERATION_IMPORT_ROWS} عملية في الملف الواحد.")

            name = _clean_text(values[0] if len(values) > 0 else None)
            description = _clean_text(values[1] if len(values) > 1 else None)
            group_name = _clean_text(values[2] if len(values) > 2 else None)

            if not name:
                raise ValueError(f"الصف {row_number}: اسم العملية مطلوب.")
            if len(name) > 120:
                raise ValueError(f"الصف {row_number}: اسم العملية لا يتجاوز 120 حرفًا.")
            name_key = _identity(name)
            if name_key in seen_names:
                raise ValueError(f"الصف {row_number}: العملية «{name}» مكررة داخل الملف.")
            seen_names.add(name_key)

            group_id = None
            if group_name and group_name != GROUP_NO_VALUE:
                group_key = _identity(group_name)
                if group_key in ambiguous_groups:
                    raise ValueError(f"الصف {row_number}: اسم المجموعة «{group_name}» غير فريد في البرنامج.")
                if group_key not in group_lookup:
                    raise ValueError(f"الصف {row_number}: المجموعة «{group_name}» غير موجودة في مكتبة البرنامج.")
                group_id = group_lookup[group_key]

            condition = _parse_fixed_axes(
                values,
                start_index=3,
                axes=(("interval_km", "كم"), ("interval_hours", "ساعة"), ("interval_days", "يوم")),
                field_name="الشرط",
                row_number=row_number,
                required=True,
            )
            warnings = _parse_fixed_axes(
                values,
                start_index=6,
                axes=(("warning_km", "كم"), ("warning_hours", "ساعة"), ("warning_days", "يوم")),
                field_name="الإنذار",
                row_number=row_number,
                required=True,
                warning=True,
            )

            rows.append({
                "name": name,
                "description": description or None,
                "group_id": group_id,
                **warnings,
                "condition": condition,
            })

        if not rows:
            raise ValueError("أضف عملية واحدة على الأقل في ورقة «العمليات».")
        return rows
    finally:
        workbook.close()


def import_model_operations(db: Session, *, model_id: int, data: bytes) -> int:
    """Create model-scoped operations atomically, including each model condition."""
    model = db.get(EquipmentModel, model_id)
    if model is None:
        raise ValueError("الطراز المحدد غير موجود.")

    groups = db.query(MaintenanceOperationGroup).order_by(
        MaintenanceOperationGroup.sort_order,
        MaintenanceOperationGroup.name,
    ).all()
    rows = parse_operation_workbook(data, groups)

    existing_names = {
        _identity(name)
        for (name,) in db.query(MaintenanceModelOperation.name).filter(
            MaintenanceModelOperation.equipment_model_id == model_id,
        ).all()
    }
    conflicts = [row["name"] for row in rows if _identity(row["name"]) in existing_names]
    if conflicts:
        names = "، ".join(conflicts[:5])
        suffix = " وغيرها" if len(conflicts) > 5 else ""
        raise ValueError(f"توجد عمليات بهذا الاسم في مكتبة الطراز مسبقًا: {names}{suffix}. لم يُستورد أي صف.")

    try:
        for row in rows:
            operation = MaintenanceModelOperation(
                equipment_model_id=model_id,
                name=row["name"],
                description=row["description"],
                group_id=row["group_id"],
                warning_km=row["warning_km"],
                warning_hours=row["warning_hours"],
                warning_days=row["warning_days"],
                interval_km=row["condition"]["interval_km"],
                interval_hours=row["condition"]["interval_hours"],
                interval_days=row["condition"]["interval_days"],
                is_active=True,
            )
            db.add(operation)
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise ValueError("تعذر حفظ العمليات. أُلغي الاستيراد بالكامل؛ راجع حالة قاعدة البيانات ثم أعد المحاولة.") from exc
    return len(rows)


def parse_general_operation_workbook(
    data: bytes,
    groups: list[MaintenanceOperationGroup],
) -> list[dict]:
    """Validate public-library rows without importing plan conditions."""
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, ValueError, zipfile.BadZipFile, OSError) as exc:
        raise ValueError("تعذر قراءة الملف. ارفع ملف Excel صالحًا بصيغة .xlsx.") from exc
    try:
        if "العمليات" not in workbook.sheetnames:
            raise ValueError("يجب أن يحتوي الملف على ورقة باسم «العمليات».")
        sheet = workbook["العمليات"]
        headers = tuple(
            _clean_text(cell)
            for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        )
        if headers != GENERAL_OPERATION_IMPORT_HEADERS:
            expected_headers = "، ".join(GENERAL_OPERATION_IMPORT_HEADERS)
            raise ValueError(f"رؤوس أعمدة المكتبة العامة أو ترتيبها غير مطابق. استخدم: {expected_headers}.")

        group_lookup: dict[str, int] = {}
        ambiguous_groups: set[str] = set()
        for group in groups:
            key = _identity(group.name)
            if key in group_lookup and group_lookup[key] != group.id:
                ambiguous_groups.add(key)
            group_lookup[key] = group.id

        rows: list[dict] = []
        seen_names: set[str] = set()
        row_count = 0
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(GENERAL_OPERATION_IMPORT_HEADERS), values_only=True),
            start=2,
        ):
            values = list(values)
            if not any(_clean_text(value) for value in values):
                continue
            row_count += 1
            if row_count > MAX_OPERATION_IMPORT_ROWS:
                raise ValueError(f"الحد الأقصى {MAX_OPERATION_IMPORT_ROWS} عملية في الملف الواحد.")

            name = _clean_text(values[0])
            description = _clean_text(values[1])
            group_name = _clean_text(values[2])
            if not name:
                raise ValueError(f"الصف {row_number}: اسم العملية مطلوب.")
            if len(name) > 120:
                raise ValueError(f"الصف {row_number}: اسم العملية لا يتجاوز 120 حرفًا.")
            name_key = _identity(name)
            if name_key in seen_names:
                raise ValueError(f"الصف {row_number}: العملية «{name}» مكررة داخل الملف.")
            seen_names.add(name_key)

            group_id = None
            if group_name and _identity(group_name) != _identity(GROUP_NO_VALUE):
                group_key = _identity(group_name)
                if group_key in ambiguous_groups:
                    raise ValueError(f"الصف {row_number}: اسم المجموعة «{group_name}» غير فريد في البرنامج.")
                if group_key not in group_lookup:
                    raise ValueError(f"الصف {row_number}: المجموعة «{group_name}» غير موجودة في مكتبة البرنامج.")
                group_id = group_lookup[group_key]

            warnings = {}
            for index, (field, label) in enumerate((
                ("warning_km", "إنذار كم"),
                ("warning_hours", "إنذار ساعة تشغيل"),
                ("warning_days", "إنذار يوم"),
            ), start=3):
                value = _nonnegative_number(
                    values[index], field=label, row_number=row_number,
                )
                if field == "warning_days" and value is not None:
                    if value != value.to_integral_value():
                        raise ValueError(f"الصف {row_number}: قيمة «إنذار يوم» يجب أن تكون عددًا صحيحًا.")
                    value = int(value)
                warnings[field] = value

            if all(value is None for value in warnings.values()):
                raise ValueError(f"الصف {row_number}: أدخل قيمة إنذار واحدة على الأقل.")
            rows.append({
                "name": name,
                "description": description or None,
                "group_id": group_id,
                "is_active": True,
                **warnings,
            })

        if not rows:
            raise ValueError("أضف عملية واحدة على الأقل في ورقة «العمليات».")
        return rows
    finally:
        workbook.close()


def import_general_operations(db: Session, *, data: bytes) -> int:
    """Create reusable public operations atomically, without any model conditions."""
    groups = db.query(MaintenanceOperationGroup).order_by(
        MaintenanceOperationGroup.sort_order,
        MaintenanceOperationGroup.name,
    ).all()
    rows = parse_general_operation_workbook(data, groups)
    existing_names = {
        _identity(name) for (name,) in db.query(MaintenanceOperation.name).all()
    }
    conflicts = [row["name"] for row in rows if _identity(row["name"]) in existing_names]
    if conflicts:
        names = "، ".join(conflicts[:5])
        suffix = " وغيرها" if len(conflicts) > 5 else ""
        raise ValueError(f"توجد عمليات بهذه الأسماء في المكتبة العامة مسبقًا: {names}{suffix}. لم يُستورد أي صف.")

    try:
        for row in rows:
            db.add(MaintenanceOperation(**row))
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise ValueError("تعذر حفظ العمليات العامة. أُلغي الاستيراد بالكامل؛ راجع حالة قاعدة البيانات ثم أعد المحاولة.") from exc
    return len(rows)
