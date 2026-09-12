from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


def link_models_to_configurations(db: Session, content: bytes) -> None:
    """Resolve configuration codes from the workbook after import and link them to models."""
    from io import BytesIO
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    if "Models" not in wb.sheetnames:
        return
    for row in wb["Models"].iter_rows(min_row=2, values_only=True):
        values = list(row)
        if not any(v is not None and str(v).strip() for v in values):
            continue
        headers = [str(v).strip() if v is not None else "" for v in next(wb["Models"].iter_rows(min_row=1, max_row=1, values_only=True))]
        data = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
        name = str(data.get("name") or "").strip()
        type_name = str(data.get("type_name") or "").strip()
        brand_name = str(data.get("brand_name") or "").strip()
        if not name or not type_name or not brand_name:
            continue
        model_id = db.execute(text("""SELECT em.id FROM equipment_models em
            JOIN equipment_types et ON et.id=em.equipment_type_id
            JOIN equipment_brands eb ON eb.id=em.brand_id
            WHERE em.name=:name AND et.name=:type_name AND eb.name=:brand_name"""), {"name": name, "type_name": type_name, "brand_name": brand_name}).scalar_one_or_none()
        if model_id is None:
            continue
        tire_code = str(data.get("tire_config_code") or "").strip() or None
        battery_code = str(data.get("battery_config_code") or "").strip() or None
        tire_id = db.execute(text("SELECT id FROM master_data_configurations WHERE code=:code AND config_type='TIRES'"), {"code": tire_code}).scalar_one_or_none() if tire_code else None
        battery_id = db.execute(text("SELECT id FROM master_data_configurations WHERE code=:code AND config_type='BATTERY'"), {"code": battery_code}).scalar_one_or_none() if battery_code else None
        db.execute(text("""UPDATE equipment_models SET tire_configuration_id=:tire_id,
            battery_configuration_id=:battery_id WHERE id=:id"""), {"tire_id": tire_id, "battery_id": battery_id, "id": model_id})
