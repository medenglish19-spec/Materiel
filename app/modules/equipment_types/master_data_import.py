from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


REQUIRED_SHEETS = {"Categories", "Types", "Brands", "Models"}


def _value(row: dict[str, Any], key: str, default: Any = None) -> Any:
    value = row.get(key)
    if value is None:
        return default
    if isinstance(value, str):
        value = value.strip()
    return default if value == "" else value


def _parse_bool(value: Any, field: str, sheet: str, row_no: int, default: bool = False) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().lower()
    if normalized in {"true", "1", "yes", "y", "نعم"}:
        return True
    if normalized in {"false", "0", "no", "n", "لا"}:
        return False
    raise ValueError(f"{sheet}: الصف {row_no}: {field} يجب أن يكون True/False")


def _rows(ws) -> list[dict[str, Any]]:
    values = list(ws.values)
    if not values:
        return []
    headers = [str(v).strip() if v is not None else "" for v in values[0]]
    return [{headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]} for row in values[1:] if any(v is not None and str(v).strip() for v in row)]


def _require(row: dict[str, Any], field: str, sheet: str, row_no: int) -> Any:
    value = _value(row, field)
    if value is None:
        raise ValueError(f"{sheet}: الصف {row_no}: الحقل {field} مطلوب")
    return value


def _int(value: Any, field: str, sheet: str, row_no: int, minimum: int | None = None) -> int | None:
    if value is None:
        return None
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{sheet}: الصف {row_no}: {field} يجب أن يكون رقمًا صحيحًا") from exc
    if minimum is not None and result < minimum:
        raise ValueError(f"{sheet}: الصف {row_no}: {field} لا يمكن أن يكون أقل من {minimum}")
    return result


def _float(value: Any, field: str, sheet: str, row_no: int) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{sheet}: الصف {row_no}: {field} يجب أن يكون رقمًا") from exc


def _upsert_by_name(db: Session, table: str, name: str, values: dict[str, Any]) -> int:
    current = db.execute(text(f"SELECT id FROM {table} WHERE name = :name"), {"name": name}).scalar_one_or_none()
    if current is None:
        columns = ["name", *values.keys()]
        params = {"name": name, **values}
        placeholders = ", ".join(f":{c}" for c in columns)
        db.execute(text(f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders})"), params)
        if db.bind.dialect.name == "sqlite":
            return int(db.execute(text("SELECT last_insert_rowid()")).scalar_one())
        return int(db.execute(text(f"SELECT id FROM {table} WHERE name = :name"), {"name": name}).scalar_one())
    if values:
        assignments = ", ".join(f"{k} = :{k}" for k in values)
        db.execute(text(f"UPDATE {table} SET {assignments} WHERE id = :id"), {**values, "id": current})
    return int(current)


def _config(db: Session, code: str, config_type: str, name: str) -> int:
    current = db.execute(text("SELECT id, config_type FROM master_data_configurations WHERE code = :code"), {"code": code}).one_or_none()
    now = datetime.now(timezone.utc)
    if current is None:
        db.execute(text("""INSERT INTO master_data_configurations
            (code, config_type, name, is_active, created_at, updated_at)
            VALUES (:code, :type, :name, 1, :now, :now)"""), {"code": code, "type": config_type, "name": name, "now": now})
        return int(db.execute(text("SELECT id FROM master_data_configurations WHERE code = :code"), {"code": code}).scalar_one())
    if current[1] != config_type:
        raise ValueError(f"Configuration {code} موجود مسبقًا بنوع {current[1]} ولا يمكن تغييره إلى {config_type}")
    db.execute(text("""UPDATE master_data_configurations SET name=:name,
        is_active=1, updated_at=:now WHERE id=:id"""), {"name": name, "now": now, "id": current[0]})
    return int(current[0])


def _config_item(db: Session, config_id: int, data: dict[str, Any]) -> None:
    now = datetime.now(timezone.utc)
    current = db.execute(text("""SELECT id FROM master_data_configuration_items
        WHERE configuration_id=:cid AND item_code=:code"""), {"cid": config_id, "code": data["item_code"]}).scalar_one_or_none()
    if current is None:
        db.execute(text("""INSERT INTO master_data_configuration_items
            (configuration_id,item_code,item_name,axle_number,side,position_type,sort_order,
             value_text,value_number,unit,extra_json,created_at,updated_at)
            VALUES (:configuration_id,:item_code,:item_name,:axle_number,:side,:position_type,
                    :sort_order,:value_text,:value_number,:unit,:extra_json,:now,:now)"""), {**data, "configuration_id": config_id, "now": now})
    else:
        db.execute(text("""UPDATE master_data_configuration_items SET item_name=:item_name,
            axle_number=:axle_number, side=:side, position_type=:position_type, sort_order=:sort_order,
            value_text=:value_text, value_number=:value_number, unit=:unit, extra_json=:extra_json,
            updated_at=:now WHERE id=:id"""), {**data, "now": now, "id": current})


def _sync_model_configuration_fields(db: Session) -> None:
    db.execute(text("""
        UPDATE equipment_models SET has_tires = 1,
            tire_positions_required = COALESCE((SELECT COUNT(*) FROM master_data_configuration_items i
            WHERE i.configuration_id = equipment_models.tire_configuration_id), 0)
        WHERE tire_configuration_id IS NOT NULL
    """))
    db.execute(text("""
        UPDATE equipment_models SET has_batteries = 1,
            battery_count_required = COALESCE((SELECT COUNT(*) FROM master_data_configuration_items i
            WHERE i.configuration_id = equipment_models.battery_configuration_id), 0)
        WHERE battery_configuration_id IS NOT NULL
    """))
    _materialize_tire_positions(db)


def _materialize_tire_positions(db: Session) -> None:
    """Materialize fixed model tire positions; never rewrite a historical position."""
    models = db.execute(text("""SELECT id, tire_configuration_id FROM equipment_models
        WHERE tire_configuration_id IS NOT NULL""")).all()
    for model_id, config_id in models:
        items = db.execute(text("""SELECT item_code, item_name, axle_number, side, position_type, sort_order
            FROM master_data_configuration_items WHERE configuration_id=:config_id
            ORDER BY sort_order, id"""), {"config_id": config_id}).all()
        for item_code, item_name, axle_number, side, position_type, sort_order in items:
            code = f"MD-{model_id}-{item_code}"[:40]
            current = db.execute(text("SELECT id, name, axle_number, side, position_type, sort_order FROM tire_positions WHERE code=:code"), {"code": code}).one_or_none()
            values = {"code": code, "name": item_name, "description": None, "sort_order": sort_order or 0,
                      "equipment_model_id": model_id, "axle_number": axle_number, "side": side, "position_type": position_type}
            if current is None:
                db.execute(text("""INSERT INTO tire_positions
                    (code,name,description,sort_order,equipment_model_id,axle_number,side,position_type)
                    VALUES (:code,:name,:description,:sort_order,:equipment_model_id,:axle_number,:side,:position_type)"""), values)
                continue
            changed = tuple(current[1:]) != (item_name, axle_number, side, position_type, sort_order or 0)
            if not changed:
                continue
            movement_count = db.execute(text("SELECT COUNT(*) FROM tire_movements WHERE position_id=:id"), {"id": current[0]}).scalar_one()
            if movement_count:
                raise ValueError(f"موضع الإطار {code} مرتبط بـ {movement_count} حركة تاريخية؛ لا يمكن تغيير تعريفه في مكانه. أنشئ Configuration/Position جديدًا للتعريف الجديد حتى تبقى الحركات السابقة كما هي.")
            db.execute(text("""UPDATE tire_positions SET name=:name, description=:description, sort_order=:sort_order,
                equipment_model_id=:equipment_model_id, axle_number=:axle_number, side=:side, position_type=:position_type
                WHERE id=:id"""), {**values, "id": current[0]})


def import_master_data(db: Session, content: bytes) -> dict[str, int]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheets = set(wb.sheetnames)
    missing = REQUIRED_SHEETS - sheets
    if missing:
        raise ValueError("الأوراق الإلزامية المفقودة: " + ", ".join(sorted(missing)))
    counts = {"categories": 0, "types": 0, "brands": 0, "models": 0, "tire_positions": 0, "battery_items": 0, "properties": 0}
    category_ids: dict[str, int] = {}
    brand_ids: dict[str, int] = {}
    model_ids: dict[tuple[str, str, str], int] = {}
    config_ids: dict[str, int] = {}
    with db.begin():
        for n, row in enumerate(_rows(wb["Categories"]), 2):
            code = str(_require(row, "code", "Categories", n)); name = str(_require(row, "name", "Categories", n))
            category_ids[code] = _upsert_by_name(db, "equipment_categories", name, {"code": code, "sort_order": _int(_value(row, "sort_order", 0), "sort_order", "Categories", n, 0) or 0})
            counts["categories"] += 1
        for n, row in enumerate(_rows(wb["Brands"]), 2):
            name = str(_require(row, "name", "Brands", n)); brand_ids[name] = _upsert_by_name(db, "equipment_brands", name, {"is_active": True}); counts["brands"] += 1
        for n, row in enumerate(_rows(wb["Types"]), 2):
            name = str(_require(row, "name", "Types", n)); unit = str(_require(row, "measurement_unit", "Types", n)); category_code = str(_require(row, "category_code", "Types", n))
            if category_code not in category_ids: raise ValueError(f"Types: الصف {n}: category_code غير موجود: {category_code}")
            _upsert_by_name(db, "equipment_types", name, {"measurement_unit": unit, "category_id": category_ids[category_code], "theoretical_quantity": _int(_value(row, "theoretical_quantity"), "theoretical_quantity", "Types", n, 0)}); counts["types"] += 1
        for sheet, config_type in (("TirePositions", "TIRES"), ("BatteryConfigurations", "BATTERY")):
            if sheet not in sheets: continue
            for n, row in enumerate(_rows(wb[sheet]), 2):
                code = str(_require(row, "config_code", sheet, n)); name = str(_require(row, "config_name", sheet, n)); cid = _config(db, code, config_type, name); config_ids[code] = cid
                _config_item(db, cid, {"item_code": str(_require(row, "item_code", sheet, n)), "item_name": str(_require(row, "item_name", sheet, n)), "axle_number": _int(_value(row, "axle_number"), "axle_number", sheet, n, 1), "side": _value(row, "side"), "position_type": _value(row, "position_type"), "sort_order": _int(_value(row, "sort_order", 0), "sort_order", sheet, n, 0) or 0, "value_text": _value(row, "value_text"), "value_number": _float(_value(row, "value_number"), "value_number", sheet, n), "unit": _value(row, "unit"), "extra_json": _value(row, "extra_json")})
                counts["tire_positions" if config_type == "TIRES" else "battery_items"] += 1
        model_values = list(wb["Models"].values)
        model_headers = {str(v).strip() for v in model_values[0] if v is not None and str(v).strip()} if model_values else set()
        has_tire_config_column = "tire_config_code" in model_headers; has_battery_config_column = "battery_config_code" in model_headers
        for n, row in enumerate(_rows(wb["Models"]), 2):
            name = str(_require(row, "name", "Models", n)); type_name = str(_require(row, "type_name", "Models", n)); brand_name = str(_require(row, "brand_name", "Models", n))
            type_id = db.execute(text("SELECT id FROM equipment_types WHERE name=:name"), {"name": type_name}).scalar_one_or_none()
            if type_id is None or brand_name not in brand_ids: raise ValueError(f"Models: الصف {n}: النوع أو العلامة غير موجودة")
            tire_code = str(_value(row, "tire_config_code") or "").strip() or None; battery_code = str(_value(row, "battery_config_code") or "").strip() or None
            if tire_code and tire_code not in config_ids: raise ValueError(f"Models: الصف {n}: tire_config_code غير موجود: {tire_code}")
            if battery_code and battery_code not in config_ids: raise ValueError(f"Models: الصف {n}: battery_config_code غير موجود: {battery_code}")
            mobility = _value(row, "mobility_type", "mobile"); requires_driver = _parse_bool(_value(row, "requires_driver", True), "requires_driver", "Models", n, True)
            existing = db.execute(text("SELECT id FROM equipment_models WHERE equipment_type_id=:tid AND brand_id=:bid AND name=:name"), {"tid": type_id, "bid": brand_ids[brand_name], "name": name}).scalar_one_or_none()
            values = {"equipment_type_id": type_id, "brand_id": brand_ids[brand_name], "mobility_type": mobility, "requires_driver": requires_driver}
            if has_tire_config_column: values["tire_configuration_id"] = config_ids.get(tire_code) if tire_code else None
            if has_battery_config_column: values["battery_configuration_id"] = config_ids.get(battery_code) if battery_code else None
            if existing is None:
                db.execute(text("""INSERT INTO equipment_models
                    (name,equipment_type_id,brand_id,has_tires,tire_positions_required,tire_size,has_batteries,battery_count_required,battery_capacity_ah,battery_voltage_v,mobility_type,requires_driver,tire_configuration_id,battery_configuration_id)
                    VALUES (:name,:equipment_type_id,:brand_id,0,0,NULL,0,0,NULL,NULL,:mobility_type,:requires_driver,:tire_configuration_id,:battery_configuration_id)"""), {"name": name, **values, "tire_configuration_id": values.get("tire_configuration_id"), "battery_configuration_id": values.get("battery_configuration_id")})
                existing = db.execute(text("SELECT id FROM equipment_models WHERE equipment_type_id=:tid AND brand_id=:bid AND name=:name"), {"tid": type_id, "bid": brand_ids[brand_name], "name": name}).scalar_one()
            else:
                db.execute(text("UPDATE equipment_models SET mobility_type=:mobility_type, requires_driver=:requires_driver WHERE id=:id"), {"mobility_type": mobility, "requires_driver": requires_driver, "id": existing})
                config_updates = {}
                if has_tire_config_column: config_updates["tire_configuration_id"] = values.get("tire_configuration_id")
                if has_battery_config_column: config_updates["battery_configuration_id"] = values.get("battery_configuration_id")
                if config_updates:
                    assignments = ", ".join(f"{k}=:{k}" for k in config_updates); db.execute(text(f"UPDATE equipment_models SET {assignments} WHERE id=:id"), {**config_updates, "id": existing})
            model_ids[(name, type_name, brand_name)] = int(existing); counts["models"] += 1
        if "ModelProperties" in sheets:
            now = datetime.now(timezone.utc)
            for n, row in enumerate(_rows(wb["ModelProperties"]), 2):
                key = (str(_require(row, "model_name", "ModelProperties", n)), str(_require(row, "type_name", "ModelProperties", n)), str(_require(row, "brand_name", "ModelProperties", n))); model_id = model_ids.get(key)
                if model_id is None: raise ValueError(f"ModelProperties: الصف {n}: الطراز غير موجود في Models")
                prop_key = str(_require(row, "property_key", "ModelProperties", n)); label = str(_require(row, "property_label", "ModelProperties", n))
                data = {"equipment_model_id": model_id, "property_key": prop_key, "property_label": label, "value_text": _value(row, "value_text"), "value_number": _float(_value(row, "value_number"), "value_number", "ModelProperties", n), "unit": _value(row, "unit"), "sort_order": _int(_value(row, "sort_order", 0), "sort_order", "ModelProperties", n, 0) or 0, "extra_json": _value(row, "extra_json")}
                existing = db.execute(text("SELECT id FROM master_data_model_properties WHERE equipment_model_id=:mid AND property_key=:key"), {"mid": model_id, "key": prop_key}).scalar_one_or_none()
                if existing is None:
                    db.execute(text("""INSERT INTO master_data_model_properties
                        (equipment_model_id,property_key,property_label,value_text,value_number,unit,sort_order,extra_json,created_at,updated_at)
                        VALUES (:equipment_model_id,:property_key,:property_label,:value_text,:value_number,:unit,:sort_order,:extra_json,:now,:now)"""), {**data, "now": now})
                else:
                    db.execute(text("""UPDATE master_data_model_properties SET property_label=:property_label,value_text=:value_text,value_number=:value_number,unit=:unit,sort_order=:sort_order,extra_json=:extra_json,updated_at=:now WHERE id=:id"""), {**data, "now": now, "id": existing})
                counts["properties"] += 1
        _sync_model_configuration_fields(db)
    return counts
