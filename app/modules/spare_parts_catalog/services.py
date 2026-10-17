from __future__ import annotations

import csv
import base64
import hashlib
import html.parser
import io
import json
import os
import re
import uuid
import zipfile
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.modules.faults_repairs.models import SparePart
from .models import SparePartCatalogDetails, SparePartCatalogImport, SparePartCatalogReference, SparePartCategory

from app.modules.document_import.upload import MAX_UPLOAD_BYTES
MAX_ROWS = 1000
MAX_ZIP_FILES = 100
MAX_ZIP_UNPACKED = 30 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".pdf", ".docx", ".xml", ".html", ".htm", ".zip"}

HEADER_NAMES = {
    "reference": {"reference", "ref", "part number", "part no", "part_number", "part code", "item code", "item no", "article code", "code article", "référence", "référence pièce", "référence fabricant", "référence constructeur", "référence fournisseur", "reférence", "ref.", "ref constructeur", "ref fabricant", "oem", "oem number", "oem reference", "p/n", "n° pièce", "n° de pièce", "numéro de pièce", "رقم القطعة", "رقم المرجع", "رقم الصنف", "كود القطعة", "المرجع", "مرجع"},
    "name": {"name", "part name", "item name", "product name", "part description", "item description", "product description", "description", "designation", "désignation", "désignation pièce", "désignation article", "libellé", "libellé article", "اسم المنتج", "اسم الصنف", "اسم القطعة", "التسمية", "وصف القطعة", "الوصف"},
    "name_ar": {"name ar", "arabic name", "nom arabe", "الاسم العربي", "اسم عربي"},
    "name_fr": {"name fr", "french name", "nom français", "nom francais", "الاسم الفرنسي", "اسم فرنسي"},
    "manufacturer": {"manufacturer", "manufacturer name", "make", "brand", "marque", "fabricant", "constructeur", "الشركة المصنعة", "اسم المصنع", "الشركة", "المصنع"},
    "category": {"category", "group", "family", "product family", "catégorie", "famille", "système", "system", "التصنيف", "الفئة", "المجموعة", "النظام"},
    "notes": {"notes", "comment", "remarks", "الملاحظات", "ملاحظات"},
}


def _norm(value: Any) -> str:
    text = str(value or "").strip().casefold()
    return re.sub(r"[\s\-_/]+", "", text)


def _header_norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").casefold())
    text = "".join(character for character in text if unicodedata.category(character) != "Mn")
    return re.sub(r"[^\w]+", "", text, flags=re.UNICODE)


HEADER_LOOKUP = {
    _header_norm(alias): field
    for field, aliases in HEADER_NAMES.items()
    for alias in aliases
}


def _header_mapping(row: list[Any]) -> dict[str, int]:
    mapping = {}
    for index, value in enumerate(row):
        field = HEADER_LOOKUP.get(_header_norm(value))
        if field and field not in mapping:
            mapping[field] = index
    return mapping


def _find_header(rows: list[tuple[int, list[Any]]]) -> tuple[int | None, dict[str, int]]:
    """Find the strongest recognized header near the top of a sheet."""
    candidates = []
    for position, (_source_index, row) in enumerate(rows[:20]):
        mapping = _header_mapping(row)
        if mapping:
            score = len(mapping) * 10 + (5 if {"name", "name_ar", "name_fr"} & mapping.keys() else 0)
            score += 2 if "reference" in mapping else 0
            candidates.append((score, -position, position, mapping))
    if not candidates:
        return None, {}
    _, _, position, mapping = max(candidates)
    return position, mapping


def _cell_text(value: Any, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and number_format and re.fullmatch(r"0+", number_format):
        return str(value).zfill(len(number_format))
    return str(value).strip()


class _TableParser(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._cell = []

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if any(self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv_dialect(text: str) -> csv.Dialect:
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        # Vendor exports often prepend a title row that does not have the same
        # number of fields as the table. Infer the separator from the most
        # consistent non-empty rows instead of silently treating them as CSV
        # comma-delimited data.
        lines = [line for line in text[:4096].splitlines() if line.strip()]
        candidates = []
        for delimiter in (";", "\t", "|", ","):
            counts = [line.count(delimiter) for line in lines]
            positive = [count for count in counts if count]
            if positive:
                mode = max(set(positive), key=positive.count)
                candidates.append((positive.count(mode), sum(positive), len(positive), delimiter))
        if not candidates:
            return csv.excel
        delimiter = max(candidates)[-1]
        dialect = csv.excel()
        dialect.delimiter = delimiter
        return dialect


def _xml_matrix(data: bytes) -> list[list[str]]:
    try:
        root = ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise ValueError("ملف XML تالف أو غير صالح.") from exc
    rows = []
    for node in root.iter():
        children = list(node)
        if len(children) >= 2 and all(not list(child) for child in children):
            names = [child.tag.rsplit("}", 1)[-1] for child in children]
            values = [(child.text or "").strip() for child in children]
            if not rows:
                rows.append(names)
            rows.append(values)
    return rows


def _xlsx_sheets(data: bytes) -> list[tuple[str, list[list[str]]]]:
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportError as exc:
        raise ValueError("استيراد XLSX غير متاح: مكتبة openpyxl غير مثبتة.") from exc
    except Exception as exc:
        raise ValueError("ملف Excel تالف أو غير صالح.") from exc
    result = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows():
                rows.append([_cell_text(cell.value, cell.number_format) for cell in row])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("تجاوزت الورقة الحد الأقصى وهو 1000 صف.")
            if any(any(value for value in row) for row in rows):
                result.append((sheet.title, rows))
    finally:
        workbook.close()
    return result


def _extract_document_tables(extension: str, data: bytes) -> list[tuple[str, list[list[str]]]]:
    from app.modules.document_import.extraction.extractor import ExtractionError, extract_document
    import tempfile

    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("امتداد الملف لا يطابق محتواه؛ الملف ليس PDF صالحًا.")
    if extension == ".xls" and not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        raise ValueError("امتداد الملف لا يطابق محتواه؛ الملف ليس XLS صالحًا.")
    if extension == ".xls":
        try:
            import xlrd  # noqa: F401
        except ImportError as exc:
            raise ValueError("ملفات XLS القديمة تحتاج إلى xlrd؛ احفظ الملف بصيغة XLSX أو CSV.") from exc
    suffix = extension if extension in {".pdf", ".docx", ".xls"} else ".pdf"
    path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temp:
            temp.write(data)
            path = temp.name
        try:
            extracted = extract_document(path)
        except ExtractionError as exc:
            raise ValueError(exc.message) from exc
        tables = [(table.name or "", table.rows) for table in extracted.tables if table.rows]
        if tables:
            return tables
        if extracted.text:
            lines = []
            for line in extracted.text.splitlines():
                cells = [cell.strip() for cell in re.split(r"\t+|\s{2,}|\s*;\s*", line) if cell.strip()]
                if len(cells) > 1:
                    lines.append(cells)
            if len(lines) > 1:
                return [("", lines)]
        raise ValueError("لم يعثر المستخرج الحالي على جدول قطع قابل للمراجعة.")
    finally:
        if path:
            try:
                os.unlink(path)
            except OSError:
                pass


def _parse_content(filename: str, data: bytes, depth: int = 0) -> tuple[list[dict], str, list[str]]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("الصيغة غير مدعومة. المتاح: CSV وXLSX وPDF وDOCX وXLS وXML وHTML وZIP.")
    if depth > 0 and extension == ".zip":
        raise ValueError("الأرشيفات المتداخلة غير مدعومة لأسباب السلامة.")
    warnings: list[str] = []
    tables: list[tuple[str, list[list[str]]]]
    if extension == ".csv":
        text = _decode_text(data)
        dialect = _csv_dialect(text)
        tables = [("", list(csv.reader(io.StringIO(text), dialect)))]
    elif extension == ".xlsx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError("امتداد الملف لا يطابق محتواه؛ الملف ليس XLSX صالحًا.")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("الملف المضغوط ليس مصنف XLSX.")
            if len(infos) > MAX_ZIP_FILES or sum(item.file_size for item in infos) > MAX_ZIP_UNPACKED:
                raise ValueError("محتوى XLSX يتجاوز الحدود الآمنة.")
            if any(item.compress_size and item.file_size / item.compress_size > 100 for item in infos):
                raise ValueError("نسبة ضغط XLSX مرتفعة جدًا.")
        tables = _xlsx_sheets(data)
    elif extension in {".html", ".htm"}:
        text = _decode_text(data)
        if "<html" not in text[:4096].casefold() and "<table" not in text[:4096].casefold():
            raise ValueError("امتداد الملف لا يطابق محتوى HTML.")
        parser = _TableParser()
        parser.feed(text)
        tables = [(f"table {index + 1}", rows) for index, rows in enumerate(parser.tables)]
    elif extension == ".xml":
        tables = [("XML", _xml_matrix(data))]
    elif extension == ".zip":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError("ملف ZIP تالف أو غير صالح.")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_FILES:
                raise ValueError("يحتوي ZIP على عدد ملفات يتجاوز الحد الآمن.")
            total = sum(item.file_size for item in infos)
            if total > MAX_ZIP_UNPACKED:
                raise ValueError("الحجم بعد فك الضغط يتجاوز 30 ميغابايت.")
            candidates = []
            image_members = []
            ignored_files = []
            for item in infos:
                normalized_name = item.filename.replace("\\", "/")
                path = Path(normalized_name)
                mode = item.external_attr >> 16
                if path.is_absolute() or path.drive or ":" in normalized_name or normalized_name.startswith("/") or ".." in path.parts or (mode & 0o170000) == 0o120000:
                    raise ValueError("رفض ZIP لاحتوائه على مسار أو رابط رمزي غير آمن.")
                if item.is_dir():
                    continue
                if item.filename.lower().endswith(".zip"):
                    raise ValueError("الأرشيفات المتداخلة غير مدعومة.")
                if item.compress_size and item.file_size / item.compress_size > 100:
                    raise ValueError("نسبة ضغط ZIP مرتفعة جدًا.")
                if path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}:
                    image_members.append(item)
                    continue
                if Path(item.filename).suffix.lower() in SUPPORTED_EXTENSIONS - {".zip"}:
                    candidates.append(item)
                else:
                    ignored_files.append(item.filename)
            if len(candidates) != 1:
                raise ValueError("يجب أن يحتوي ZIP على ملف كتالوج مدعوم واحد بالضبط.")
            item = candidates[0]
            records, inner_type, inner_warnings = _parse_content(item.filename, archive.read(item), depth + 1)
            warnings.extend(inner_warnings)
            if ignored_files:
                warnings.append("تُركت ملفات ZIP غير المدعومة دون معالجة: " + ", ".join(Path(name).name for name in ignored_files[:5]))
            if image_members:
                if len(image_members) > 50 or sum(image.file_size for image in image_members) > 5 * 1024 * 1024:
                    raise ValueError("صور ZIP تتجاوز حد 50 صورة أو 5 ميغابايت إجمالًا.")
                by_reference = {}
                for image in image_members:
                    if image.file_size > 2 * 1024 * 1024:
                        raise ValueError("إحدى صور ZIP تتجاوز 2 ميغابايت.")
                    by_reference.setdefault(_norm(Path(image.filename).stem), []).append(image)
                linked = set()
                for record in records:
                    matches = by_reference.get(_norm(record.get("reference")), []) if record.get("reference") else []
                    if len(matches) == 1:
                        member = matches[0]
                        normalized, fmt = normalize_image_bytes(archive.read(member), Path(member.filename).suffix.lower())
                        record["image_name"] = Path(member.filename).name
                        record["image_data"] = base64.b64encode(normalized).decode("ascii")
                        record["image_format"] = fmt.lower()
                        linked.add(member.filename)
                    elif len(matches) > 1:
                        warnings.append(f"تعذر ربط الصورة بالمرجع {record['reference']}: يوجد أكثر من تطابق.")
                if len(linked) < len(image_members):
                    warnings.append("تُركت بعض صور ZIP غير المرتبطة بمرجع مطابق دون ربط.")
            return records, extension, warnings
    else:
        if extension == ".pdf" and not data.startswith(b"%PDF-"):
            raise ValueError("امتداد الملف لا يطابق محتواه؛ الملف ليس PDF صالحًا.")
        if extension == ".docx":
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise ValueError("امتداد الملف لا يطابق محتوى DOCX.")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("الملف المضغوط لا يحتوي مستند DOCX.")
            tables = _extract_document_tables(extension, data)
        elif extension == ".xls":
            if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                raise ValueError("امتداد الملف لا يطابق محتواه؛ الملف ليس XLS صالحًا.")
            try:
                import xlrd
                workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            except ImportError as exc:
                raise ValueError("ملفات XLS القديمة تحتاج إلى xlrd؛ احفظ الملف بصيغة XLSX أو CSV.") from exc
            except Exception as exc:
                raise ValueError("ملف XLS تالف أو غير صالح.") from exc
            try:
                tables = []
                for sheet in workbook.sheets():
                    if sheet.nrows > MAX_ROWS + 1:
                        raise ValueError("تجاوزت ورقة XLS الحد الأقصى وهو 1000 صف.")
                    matrix = [[_cell_text(sheet.cell_value(row, col)) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
                    if matrix:
                        tables.append((sheet.name, matrix))
            finally:
                workbook.release_resources()
        else:
            tables = _extract_document_tables(extension, data)

    records = []
    row_number = 0
    for sheet, matrix in tables:
        nonempty = [
            (index, row)
            for index, row in enumerate(matrix)
            if any(str(value or "").strip() for value in row)
        ]
        if not nonempty:
            continue
        header_position, mapping = _find_header(nonempty)
        if header_position is not None:
            start = header_position + 1
            skipped_rows = nonempty[header_position][0]
            if skipped_rows:
                warnings.append(f"الورقة {sheet or 'الأولى'}: تم تجاوز {skipped_rows} صفوف تمهيدية قبل صف العناوين.")
        else:
            # Headerless two-column files use reference, name by convention.
            start = 0
            mapping = {"reference": 0, "name": 1} if len(nonempty[0][1]) >= 2 else {}
            if mapping:
                warnings.append(f"الورقة {sheet or 'الأولى'}: لم تُكتشف عناوين؛ اعتُمد العمودان الأولان مرجعًا واسمًا. راجع المعاينة.")
        if not mapping:
            continue
        for source_index, row in nonempty[start:]:
            if header_position is not None and _header_mapping(row) == mapping:
                continue
            row_number += 1
            if row_number > MAX_ROWS:
                raise ValueError("تجاوز الملف الحد الأقصى وهو 1000 صف.")
            get = lambda field: str(row[mapping[field]] or "").strip() if field in mapping and mapping[field] < len(row) else ""
            record = {
                "source_row": source_index + 1,
                "sheet": sheet,
                "reference": get("reference")[:120],
                "name": get("name")[:160],
                "name_ar": get("name_ar")[:160],
                "name_fr": get("name_fr")[:160],
                "manufacturer": get("manufacturer")[:160],
                "category": get("category")[:255],
                "notes": get("notes")[:2000],
            }
            if not any(record[key] for key in ("reference", "name", "name_ar", "name_fr", "manufacturer")):
                continue
            if not record["name"]:
                record["name"] = record["name_fr"] or record["name_ar"]
            record["state"] = "ready" if record["name"] else "invalid"
            record["issue"] = "" if record["name"] else "اسم القطعة مطلوب."
            records.append(record)
    if not records:
        raise ValueError("لم يُعثر على صفوف قطع. تحقق من صف العناوين ووجود عمود للاسم.")
    return records, extension, warnings


def _duplicate_state(db: Session, record: dict) -> tuple[str, str]:
    ref = _norm(record.get("reference"))
    maker = _norm(record.get("manufacturer"))
    if ref:
        query = db.query(SparePartCatalogReference).filter(SparePartCatalogReference.reference == record["reference"])
        if maker:
            query = query.filter(SparePartCatalogReference.manufacturer == record["manufacturer"])
        if query.first():
            return "exact", "مرجع مطابق موجود للشركة نفسها؛ لم يتم الدمج تلقائيًا."
        if db.query(SparePart.id).filter(SparePart.part_number == record["reference"]).first():
            return "exact", "رقم القطعة موجود بالفعل؛ راجع السجل قبل الاعتماد."
    name = _norm(record.get("name"))
    if name and db.query(SparePart.id).filter(SparePart.name.ilike(record["name"])).first():
        return "possible", "يوجد اسم مطابق؛ تأكد يدويًا من اختلاف القطعة قبل الإضافة."
    return "none", ""


def create_import_preview(db: Session, filename: str, data: bytes, user_id: int | None = None) -> SparePartCatalogImport:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("حجم الملف يتجاوز 10 ميغابايت.")
    clean_name = Path(filename or "catalog").name[:255]
    records, extension, warnings = _parse_content(clean_name, data)
    digest = hashlib.sha256(data).hexdigest()
    counts = {"valid": 0, "invalid": 0, "exact": 0, "possible": 0}
    for record in records:
        duplicate, reason = _duplicate_state(db, record)
        record["duplicate"] = duplicate
        if reason and not record["issue"]:
            record["issue"] = reason
        counts["invalid" if record["state"] == "invalid" else "valid"] += 1
        counts[duplicate] = counts.get(duplicate, 0) + 1
    preview = {"rows": records, "warnings": warnings}
    stored_rows = [{key: value for key, value in row.items() if key != "image_data"} for row in records]
    payload = {"rows": stored_rows, "warnings": warnings}
    previous = db.query(SparePartCatalogImport).filter_by(content_sha256=digest).order_by(SparePartCatalogImport.id.desc()).first()
    if previous:
        if previous.status != "completed":
            previous.payload = json.dumps(payload, ensure_ascii=False)
        previous._preview_payload = preview
        return previous
    batch = SparePartCatalogImport(
        filename=clean_name,
        content_sha256=digest,
        file_type=extension.lstrip("."),
        payload=json.dumps(payload, ensure_ascii=False),
        status="preview",
        created_by=user_id,
    )
    db.add(batch)
    db.flush()
    batch._preview_payload = preview
    return batch


def preview_data(batch: SparePartCatalogImport, *, include_images: bool = False) -> dict:
    if include_images and hasattr(batch, "_preview_payload"):
        return batch._preview_payload
    return json.loads(batch.payload or "{}")


def _category_for_path(db: Session, path: str) -> SparePartCategory | None:
    parent = None
    for name in [piece.strip() for piece in re.split(r"\s*(?:>|/|\\)\s*", path or "") if piece.strip()][:8]:
        query = db.query(SparePartCategory).filter(SparePartCategory.name_ar == name)
        query = query.filter(SparePartCategory.parent_id == (parent.id if parent else None))
        category = query.first()
        if not category:
            category = SparePartCategory(name_ar=name, parent_id=parent.id if parent else None)
            db.add(category)
            db.flush()
        parent = category
    return parent


def apply_import(db: Session, batch_id: int, form, user_id: int | None = None) -> dict:
    batch = db.query(SparePartCatalogImport).filter(SparePartCatalogImport.id == batch_id).with_for_update().first()
    if not batch:
        raise ValueError("جلسة الاستيراد غير موجودة.")
    if batch.status == "completed":
        return json.loads(batch.summary or "{}")
    payload = preview_data(batch)
    source = batch.filename
    results = {"added": 0, "skipped": 0, "errors": []}
    new_images = []
    try:
        for index, original in enumerate(payload.get("rows", [])):
            if form.get(f"row-{index}-include") not in {"1", "on", "true"}:
                results["skipped"] += 1
                continue
            row = dict(original)
            for field, max_length in (("name", 160), ("name_ar", 160), ("name_fr", 160), ("manufacturer", 160), ("category", 255), ("notes", 2000), ("reference", 120)):
                value = form.get(f"row-{index}-{field}")
                if value is not None:
                    row[field] = str(value).strip()[:max_length]
            name = row.get("name") or row.get("name_fr") or row.get("name_ar")
            if not name:
                results["errors"].append({"row": index + 1, "reason": "اسم القطعة مطلوب."})
                continue
            duplicate, reason = _duplicate_state(db, row)
            if duplicate != "none" and form.get(f"row-{index}-allow-duplicate") not in {"1", "on", "true"}:
                results["skipped"] += 1
                results["errors"].append({"row": index + 1, "reason": reason})
                continue
            internal_number = "CAT-" + uuid.uuid4().hex[:16].upper()
            part = SparePart(part_number=internal_number, name=name, receiving_document=None, notes=row.get("notes") or None)
            db.add(part)
            db.flush()
            category = _category_for_path(db, row.get("category", ""))
            image_path = None
            row["image_data"] = form.get(f"row-{index}-image-data")
            if row.get("image_data"):
                try:
                    image_data = base64.b64decode(row["image_data"], validate=True)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f"بيانات صورة الصف {index + 1} غير صالحة.") from exc
                image_path = store_image(image_data, row.get("image_name", "image.png"))
                new_images.append(image_path)
            details = SparePartCatalogDetails(
                spare_part_id=part.id,
                category_id=category.id if category else None,
                name_ar=row.get("name_ar") or None,
                name_fr=row.get("name_fr") or None,
                original_name=row.get("name") or None,
                manufacturer=row.get("manufacturer") or None,
                source=source,
                image_path=image_path,
            )
            db.add(details)
            if row.get("reference"):
                db.add(SparePartCatalogReference(
                    spare_part_id=part.id,
                    manufacturer=row.get("manufacturer") or None,
                    reference=row["reference"],
                    source=source,
                ))
            results["added"] += 1
        batch.status = "completed"
        batch.completed_at = datetime.now(timezone.utc)
        batch.summary = json.dumps(results, ensure_ascii=False)
        db.commit()
    except Exception:
        db.rollback()
        for image_path in new_images:
            remove_image(image_path)
        raise
    return results


def normalize_image_bytes(data: bytes, extension: str) -> tuple[bytes, str]:
    if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
        raise ValueError("نوع الصورة غير مدعوم. استخدم JPEG أو PNG أو WebP.")
    try:
        from PIL import Image
        source = io.BytesIO(data)
        with Image.open(source) as image:
            if image.width * image.height > 50_000_000:
                raise ValueError("أبعاد الصورة تتجاوز الحد الآمن.")
            image.verify()
        source.seek(0)
        with Image.open(source) as image:
            output = io.BytesIO()
            image_format = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".webp": "WEBP"}[extension]
            if image_format == "JPEG":
                image.convert("RGB").save(output, format="JPEG", quality=90, optimize=True)
            else:
                image.save(output, format=image_format)
            return output.getvalue(), image_format
    except ImportError as exc:
        raise ValueError("رفع الصور يتطلب مكتبة Pillow المثبتة مع التطبيق.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("ملف الصورة تالف أو غير صالح.") from exc


def store_image(data: bytes, filename: str) -> str:
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("يجب ألا يتجاوز حجم الصورة 5 ميغابايت.")
    extension = Path(filename or "").suffix.lower()
    if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
        raise ValueError("الصورة غير صالحة أو نوعها غير مدعوم. استخدم JPEG أو PNG أو WebP.")
    data, normalized_format = normalize_image_bytes(data, extension)
    extension = ".jpg" if normalized_format == "JPEG" else "." + normalized_format.lower()
    root = Path(__file__).resolve().parents[3] / "static" / "uploads" / "spare-parts"
    root.mkdir(parents=True, exist_ok=True)
    safe_name = uuid.uuid4().hex + extension
    target = root / safe_name
    target.write_bytes(data)
    return f"/static/uploads/spare-parts/{safe_name}"


def remove_image(image_path: str | None) -> None:
    """Remove a catalog-owned image after its database reference is committed away."""
    prefix = "/static/uploads/spare-parts/"
    if not image_path or not image_path.startswith(prefix):
        return
    root = (Path(__file__).resolve().parents[3] / "static" / "uploads" / "spare-parts").resolve()
    target = (root / Path(image_path[len(prefix):]).name).resolve()
    if target.parent == root and target.exists() and target.is_file():
        target.unlink()
