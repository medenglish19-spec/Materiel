from dataclasses import dataclass,field
from pathlib import Path
import io
import shutil
@dataclass
class ExtractedTable:
 rows:list[list[str]]=field(default_factory=list); name:str|None=None; page:int|None=None
@dataclass
class ExtractedDocument:
 text:str=""; tables:list[ExtractedTable]=field(default_factory=list); metadata:dict=field(default_factory=dict); ocr_used:bool=False; ocr_engine:str|None=None
MAX_FILE_BYTES=15*1024*1024
def extract_document(filename,content):
 ext=Path(filename or "").suffix.lower()
 if ext not in {".pdf",".docx",".xlsx"}: raise ValueError("الصيغ المسموحة هي PDF وDOCX وXLSX")
 if not content: raise ValueError("الملف فارغ")
 if len(content)>MAX_FILE_BYTES: raise ValueError("حجم الملف يتجاوز 15 ميغابايت")
 if ext==".pdf":
  try:
   from pypdf import PdfReader
   reader=PdfReader(io.BytesIO(content),strict=False); pages=[f"[صفحة {i}]\n{p.extract_text() or ''}" for i,p in enumerate(reader.pages,1)]
   text="\n".join(p for p in pages if p.strip())
   if not text.strip():
    if shutil.which("tesseract") and shutil.which("pdftoppm"):
     try:
      from pdf2image import convert_from_bytes
      import pytesseract
      images=convert_from_bytes(content,dpi=160,fmt="png")
      text="\\n".join(f"[صفحة {i}]\\n{pytesseract.image_to_string(image)}" for i,image in enumerate(images,1))
      if text.strip():return ExtractedDocument(text=text,metadata={"pages":len(reader.pages)},ocr_used=True,ocr_engine="tesseract")
     except ImportError: pass
    raise ValueError("لم يُستخرج نص من PDF؛ يبدو ممسوحًا ضوئيًا. يلزم تثبيت Tesseract وPoppler محليًا مع pytesseract وpdf2image؛ لم يتم تنزيل أي شيء.")
   return ExtractedDocument(text=text,metadata={"pages":len(reader.pages)})
  except ValueError: raise
  except Exception as e: raise ValueError("تعذر قراءة ملف PDF") from e
 if ext==".docx":
  try:
   from docx import Document
   d=Document(io.BytesIO(content))
   tables=[ExtractedTable([[c.text.strip() for c in row.cells] for row in t.rows],f"table:{i}") for i,t in enumerate(d.tables,1)]
   return ExtractedDocument("\n".join(p.text.strip() for p in d.paragraphs if p.text.strip()),tables,{"tables":len(tables)})
  except Exception as e: raise ValueError("تعذر قراءة ملف DOCX") from e
 try:
  from openpyxl import load_workbook
  w=load_workbook(io.BytesIO(content),read_only=True,data_only=True);tables=[];parts=[]
  for s in w.worksheets:
   rows=[]
   for row in s.iter_rows(values_only=True):
    vals=[str(v).strip() if v is not None else "" for v in row]
    if any(vals):rows.append(vals);parts.append("\t".join(vals))
   tables.append(ExtractedTable(rows,s.title))
  w.close();return ExtractedDocument("\n".join(parts),tables,{"sheets":len(tables)})
 except Exception as e: raise ValueError("تعذر قراءة ملف XLSX") from e
